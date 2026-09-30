"""Bitácora de la cuenta: solo el titular la consulta, la verifica y la descarga.

- `GET /api/bitacora/?empresa=&desde=&hasta=&persona=&pagina=`: registros, del más
  reciente al más antiguo, más la lista de personas para filtrar.
- `GET /api/bitacora/verificar/`: recorre la cadena de huellas.
- `GET /api/bitacora/exportar/?formato=pdf|xlsx&…`: copia para presentar (p. ej. a
  la DT) con folio y código de verificación; la verificación pública
  (`/verificar/<código>`) confirma que salió de Jornada40 y que esos registros
  siguen iguales en el sistema.
"""
import datetime
import hashlib
import io
import secrets

from django.db import transaction
from django.db.models import Max
from django.http import HttpResponse
from django.template.loader import get_template
from django.utils import timezone
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from ..autenticacion import actor
from ..bitacora import _nombre, cuenta_de, verificar
from ..models import Empresa, ExportacionBitacora, RegistroBitacora
from ..permisos import ANIOS_BITACORA
from .base import _html_a_pdf_bytes

POR_PAGINA = 100
# Un PDF con más filas no se lee ni se imprime: se pide acotar las fechas (el Excel lleva más).
MAX_PDF = 3000
MAX_XLSX = 100000
_ALFABETO = 'ABCDEFGHJKLMNPQRSTUVWXYZ23456789'
_ACTOR = {'TITULAR': 'Titular', 'EQUIPO': 'Equipo', 'SISTEMA': 'Sistema'}


def _solo_titular(request):
    cuenta, tipo = cuenta_de(request.user)
    return cuenta if tipo == 'TITULAR' else None


def _sin_permiso():
    return Response({'error': 'Solo el titular de la cuenta puede ver la bitácora.'}, status=403)


def _filtros(request, cuenta):
    """(queryset, filtros) según la consulta; lanza ValueError si un filtro no es válido."""
    qs = RegistroBitacora.objects.filter(cuenta=cuenta)
    empresa = None
    if request.query_params.get('empresa'):
        empresa = Empresa.objects.filter(pk=int(request.query_params['empresa']), owner=cuenta).first()
        if empresa is None:
            raise ValueError('empresa')
        qs = qs.filter(empresa=empresa)
    desde = hasta = None
    if request.query_params.get('desde'):
        desde = datetime.date.fromisoformat(request.query_params['desde'])
        qs = qs.filter(creado_en__date__gte=desde)
    if request.query_params.get('hasta'):
        hasta = datetime.date.fromisoformat(request.query_params['hasta'])
        qs = qs.filter(creado_en__date__lte=hasta)
    persona = (request.query_params.get('persona') or '').strip()[:15]
    if persona:
        qs = qs.filter(actor_rut=persona)
    return qs, {'empresa': empresa, 'desde': desde, 'hasta': hasta, 'persona': persona}


def _fila(r):
    return {
        'id': r.id, 'fecha': timezone.localtime(r.creado_en).isoformat(), 'actor': r.actor_nombre,
        'actor_tipo': r.actor_tipo, 'accion': r.accion, 'descripcion': r.descripcion,
        'empresa': r.empresa.nombre_legal if r.empresa else None, 'ip': r.ip, 'resultado': r.estado_http,
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def bitacora(request):
    """Registros de la cuenta, del más reciente al más antiguo."""
    cuenta = _solo_titular(request)
    if cuenta is None:
        return _sin_permiso()
    try:
        qs, _ = _filtros(request, cuenta)
        pagina = max(int(request.query_params.get('pagina') or 1), 1)
    except ValueError:
        return Response({'error': 'Filtros inválidos.'}, status=400)
    qs = qs.select_related('empresa').order_by('-id')
    total = qs.count()
    personas = (RegistroBitacora.objects.filter(cuenta=cuenta).exclude(actor_rut='')
                .values('actor_rut').annotate(ultimo=Max('id')).order_by('-ultimo'))
    nombres = dict(RegistroBitacora.objects.filter(id__in=[p['ultimo'] for p in personas])
                   .values_list('actor_rut', 'actor_nombre'))
    return Response({
        'total': total, 'pagina': pagina, 'por_pagina': POR_PAGINA, 'anios_conservacion': ANIOS_BITACORA,
        'personas': [{'rut': p['actor_rut'], 'nombre': nombres.get(p['actor_rut'], p['actor_rut'])} for p in personas],
        'registros': [_fila(r) for r in qs[(pagina - 1) * POR_PAGINA: pagina * POR_PAGINA]],
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def verificar_bitacora(request):
    """Recorre la cadena de huellas: si alguien alteró o quitó un registro, lo informa."""
    cuenta = _solo_titular(request)
    if cuenta is None:
        return _sin_permiso()
    return Response(verificar(cuenta))


# ── Exportación verificable ──────────────────────────────────────────────────

def huella_de(hashes):
    """SHA-256 de las huellas de los registros exportados, en orden."""
    return hashlib.sha256('\n'.join(hashes).encode()).hexdigest()


def _nuevo_codigo():
    crudo = ''.join(secrets.choice(_ALFABETO) for _ in range(12))
    return f'{crudo[:4]}-{crudo[4:8]}-{crudo[8:]}'


def _titular_de(cuenta, empresa=None):
    """(nombre, rut) de quien presenta la bitácora: la empresa filtrada o el titular."""
    if empresa is not None:
        return empresa.nombre_legal, empresa.rut
    cliente = cuenta.perfil_cliente
    nombre = cliente.razon_social or ' '.join(x for x in (cliente.nombres, cliente.apellido_paterno,
                                                          cliente.apellido_materno) if x)
    return nombre, cliente.rut


def _periodo(e):
    if e.desde and e.hasta:
        return f'{e.desde:%d-%m-%Y} al {e.hasta:%d-%m-%Y}'
    if e.desde:
        return f'desde el {e.desde:%d-%m-%Y}'
    if e.hasta:
        return f'hasta el {e.hasta:%d-%m-%Y}'
    return 'todos los registros'


def _crear_exportacion(request, cuenta, qs, filtros, formato):
    ids_hashes = list(qs.order_by('id').values_list('id', 'hash'))
    persona = actor(request)
    with transaction.atomic():
        # Bloquea la cuenta: dos descargas simultáneas no toman el mismo número.
        type(cuenta).objects.select_for_update().filter(pk=cuenta.pk).first()
        numero = (ExportacionBitacora.objects.filter(cuenta=cuenta).aggregate(n=Max('numero'))['n'] or 0) + 1
        codigo = _nuevo_codigo()
        while ExportacionBitacora.objects.filter(codigo=codigo).exists():
            codigo = _nuevo_codigo()
        return ExportacionBitacora.objects.create(
            cuenta=cuenta, numero=numero, codigo=codigo, formato=formato, empresa=filtros['empresa'],
            desde=filtros['desde'], hasta=filtros['hasta'], persona=filtros['persona'],
            registros=len(ids_hashes), primer_registro=ids_hashes[0][0] if ids_hashes else None,
            ultimo_registro=ids_hashes[-1][0] if ids_hashes else None, huella=huella_de([h for _, h in ids_hashes]),
            generado_por=persona, generado_por_nombre=_nombre(persona)[:150])


def _registros_de(e):
    """Los mismos registros que tomó la exportación (la bitácora solo crece, así que el rango los fija)."""
    qs = RegistroBitacora.objects.filter(cuenta=e.cuenta)
    if e.primer_registro is None:
        return qs.none()
    qs = qs.filter(id__gte=e.primer_registro, id__lte=e.ultimo_registro)
    if e.empresa_id:
        qs = qs.filter(empresa_id=e.empresa_id)
    if e.desde:
        qs = qs.filter(creado_en__date__gte=e.desde)
    if e.hasta:
        qs = qs.filter(creado_en__date__lte=e.hasta)
    if e.persona:
        qs = qs.filter(actor_rut=e.persona)
    return qs.order_by('id')


def estado_exportacion(e):
    """Si los registros de la copia siguen iguales en el sistema."""
    return huella_de(list(_registros_de(e).values_list('hash', flat=True))) == e.huella


def verificacion_exportacion(e):
    """Respuesta de la verificación pública, con la misma forma que la de un certificado."""
    nombre, rut = _titular_de(e.cuenta, e.empresa)
    integra = estado_exportacion(e)
    return {
        'valido': integra, 'anulado': False, 'alterado': not integra, 'folio': e.folio, 'codigo': e.codigo,
        'titulo': 'Bitácora de acciones', 'emitido': f'{timezone.localtime(e.creado_en):%d-%m-%Y}',
        'empresa': {'nombre': nombre, 'rut': rut},
        'filas': [['Período', _periodo(e)], ['Registros incluidos', str(e.registros)],
                  ['Formato', e.get_formato_display()], ['Descargada por', e.generado_por_nombre],
                  ['Huella SHA-256', e.huella],
                  ['Estado', 'Los registros siguen idénticos en Jornada40.' if integra
                   else 'Los registros ya no coinciden con esta copia: fue alterada o la bitácora se modificó.']],
        'tabla': None,
        'nota': 'Cada registro guarda la huella del anterior: cambiar o quitar uno rompe la cadena y se detecta.',
    }


def _pdf(e, registros, cadena):
    from .certificados import _qr_png, url_verificacion
    nombre, rut = _titular_de(e.cuenta, e.empresa)
    url = url_verificacion(e.codigo)
    html = get_template('bitacora.html').render({
        'e': e, 'nombre': nombre, 'rut': rut, 'periodo': _periodo(e), 'url': url, 'qr': _qr_png(url),
        'generado': f'{timezone.localtime(e.creado_en):%d-%m-%Y %H:%M}', 'cadena': cadena,
        'filas': [{**_fila(r), 'fecha': f'{timezone.localtime(r.creado_en):%d-%m-%Y %H:%M:%S}',
                   'actor_tipo': _ACTOR.get(r.actor_tipo, r.actor_tipo)} for r in registros],
        'anios': ANIOS_BITACORA,
    })
    return _html_a_pdf_bytes(html, e.folio)


def _xlsx(e, registros, cadena):
    import openpyxl
    from openpyxl.styles import Font
    from .certificados import url_verificacion
    nombre, rut = _titular_de(e.cuenta, e.empresa)
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = 'Bitácora'
    cabecera = [
        ('Bitácora de acciones', f'Folio {e.folio}'), ('Titular', f'{nombre} · RUT {rut}'),
        ('Período', _periodo(e)), ('Registros', e.registros),
        ('Descargada', f'{timezone.localtime(e.creado_en):%d-%m-%Y %H:%M} por {e.generado_por_nombre}'),
        ('Código de verificación', e.codigo), ('Verificar en', url_verificacion(e.codigo)),
        ('Huella SHA-256', e.huella),
        ('Cadena de la cuenta', 'Íntegra' if cadena['ok'] else f'Alterada desde el registro {cadena["roto_en"]}'),
    ]
    for fila in cabecera:
        hoja.append(list(fila))
        hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True)
    hoja.append([])
    columnas = ['N°', 'Fecha y hora', 'Persona', 'Tipo', 'Acción', 'Detalle', 'Empresa', 'IP', 'Resultado', 'Huella']
    hoja.append(columnas)
    for c in range(1, len(columnas) + 1):
        hoja.cell(row=hoja.max_row, column=c).font = Font(bold=True)
    for r in registros:
        hoja.append([r.id, f'{timezone.localtime(r.creado_en):%d-%m-%Y %H:%M:%S}', r.actor_nombre,
                     _ACTOR.get(r.actor_tipo, r.actor_tipo), r.accion, r.descripcion,
                     r.empresa.nombre_legal if r.empresa else '', r.ip, r.estado_http or '', r.hash])
    for letra, ancho in zip('ABCDEFGHIJ', (8, 20, 30, 10, 14, 60, 30, 16, 10, 66)):
        hoja.column_dimensions[letra].width = ancho
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def exportar_bitacora(request):
    """Descarga la bitácora filtrada en PDF o Excel, con folio y código verificable."""
    cuenta = _solo_titular(request)
    if cuenta is None:
        return _sin_permiso()
    formato = (request.query_params.get('formato') or 'pdf').lower()
    if formato not in ('pdf', 'xlsx'):
        return Response({'error': 'Formato inválido.'}, status=400)
    try:
        qs, filtros = _filtros(request, cuenta)
    except ValueError:
        return Response({'error': 'Filtros inválidos.'}, status=400)
    total = qs.count()
    if total == 0:
        return Response({'error': 'No hay registros con esos filtros.'}, status=400)
    maximo = MAX_PDF if formato == 'pdf' else MAX_XLSX
    if total > maximo:
        return Response({'error': f'Son {total} registros y el {formato.upper() if formato == "pdf" else "Excel"} '
                                  f'admite hasta {maximo}. Acota las fechas'
                                  + (' o descárgala en Excel.' if formato == 'pdf' else '.')}, status=400)
    e = _crear_exportacion(request, cuenta, qs, filtros, formato.upper())
    registros = list(_registros_de(e).select_related('empresa'))
    cadena = verificar(cuenta)
    if formato == 'pdf':
        contenido, tipo = _pdf(e, registros, cadena), 'application/pdf'
    else:
        contenido, tipo = _xlsx(e, registros, cadena), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    respuesta = HttpResponse(contenido, content_type=tipo)
    respuesta['Content-Disposition'] = f'attachment; filename="bitacora_{e.folio}.{formato}"'
    respuesta['X-Codigo-Verificacion'] = e.codigo
    return respuesta
