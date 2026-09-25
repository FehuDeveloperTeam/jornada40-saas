"""Liquidaciones: emisión, PDF, libro, consolidado y exportaciones."""
from rest_framework.decorators import action
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from rest_framework import viewsets
from rest_framework.exceptions import ValidationError
from django.db import IntegrityError
from django.http import HttpResponse
from django.template.loader import get_template
from ..models import Empresa, Empleado, Contrato, Liquidacion, SolicitudFirma
from xhtml2pdf import pisa
from django.utils import timezone
import io
import zipfile
from ..indicadores import obtener_uf
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from ..serializers import LiquidacionSerializer

from .base import _es_plan_semilla, _plan_permite, logger, pdf_firmado, respuesta_pdf
from .calculo_liquidacion import _calcular_liquidacion, _pdf_liquidacion, _terminos_congelados, _terminos_vigentes, _validar_conceptos
from .parametros import _anios_de_servicio, _parametros_previsionales, _tasas_afc, _tope_en_pesos
from .previred import _linea_previred, _tasa_accidentes


class LiquidacionViewSet(viewsets.ModelViewSet):
    queryset = Liquidacion.objects.all().order_by('-anio', '-mes')
    serializer_class = LiquidacionSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        # Solo liquidaciones de empleados que pertenecen al usuario autenticado
        queryset = Liquidacion.objects.filter(
            empleado__empresa__owner=self.request.user
        ).order_by('-anio', '-mes')
        empleado_id = self.request.query_params.get('empleado', None)
        if empleado_id is not None:
            queryset = queryset.filter(empleado_id=empleado_id)
        # Filtros del proceso mensual: una empresa y un período.
        params = self.request.query_params
        if params.get('empresa'):
            queryset = queryset.filter(empleado__empresa_id=params['empresa'])
        if params.get('mes'):
            queryset = queryset.filter(mes=params['mes'])
        if params.get('anio'):
            queryset = queryset.filter(anio=params['anio'])
        return queryset

    @action(detail=False, methods=['post'])
    def simular(self, request):
        """Calcula una liquidación sin guardarla (vista previa del formulario).

        Usa exactamente el mismo cálculo que create/update: el frontend no
        replica ninguna regla previsional. Con `liquidacion` se simula la
        edición de una existente, con los términos con que se emitió.
        """
        data = request.data
        try:
            _validar_conceptos(data, request.user)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)
        try:
            empleado = Empleado.objects.get(id=data.get('empleado'), empresa__owner=request.user)
        except (Empleado.DoesNotExist, ValueError, TypeError):
            return Response({'error': 'Trabajador no encontrado o no autorizado.'}, status=status.HTTP_404_NOT_FOUND)
        contrato = Contrato.objects.filter(empleado=empleado).first()
        if not contrato:
            return Response({'error': 'El trabajador no tiene un contrato activo.'}, status=status.HTTP_400_BAD_REQUEST)

        terminos = _terminos_vigentes(contrato)
        if data.get('liquidacion'):
            existente = Liquidacion.objects.filter(id=data['liquidacion'], empleado=empleado).first()
            if existente:
                terminos = _terminos_congelados(existente, contrato)
        try:
            calculado = _calcular_liquidacion(contrato, empleado, data, terminos=terminos)
        except (ValueError, TypeError) as e:
            return Response({'error': f'Datos inválidos: {e}'}, status=status.HTTP_400_BAD_REQUEST)
        return Response({'empleado': empleado.id, 'mes': data.get('mes'), 'anio': data.get('anio'), **calculado})

    def create(self, request, *args, **kwargs):
        data = request.data
        empleado_id = data.get('empleado')

        try:
            _validar_conceptos(data, request.user)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)

        try:
            # Validar que el empleado pertenezca al usuario autenticado
            empleado = Empleado.objects.get(id=empleado_id, empresa__owner=request.user)
            contrato = Contrato.objects.filter(empleado=empleado).first()

            if not contrato:
                return Response({'error': 'El trabajador no tiene un contrato activo.'}, status=status.HTTP_400_BAD_REQUEST)

            calculado = _calcular_liquidacion(
                contrato, empleado, data, terminos=_terminos_vigentes(contrato)
            )

            liquidacion = Liquidacion.objects.create(
                empleado=empleado, mes=data.get('mes'), anio=data.get('anio'),
                **calculado,
            )

            serializer = self.get_serializer(liquidacion)
            return Response(serializer.data, status=status.HTTP_201_CREATED)

        except Empleado.DoesNotExist:
            return Response({'error': 'Trabajador no encontrado o no autorizado.'}, status=status.HTTP_404_NOT_FOUND)
        except IntegrityError:
            return Response(
                {'error': 'Ya existe una liquidación para este trabajador en el período indicado.'},
                status=status.HTTP_409_CONFLICT
            )
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    def update(self, request, *args, **kwargs):
        instance = self.get_object()

        # Una liquidación ya firmada por el trabajador queda inmutable.
        # Mientras la firma esté solo PENDIENTE, se debe cancelar esa
        # solicitud antes de poder editar (flujo ya soportado por /cancelar/).
        firmada = SolicitudFirma.objects.filter(
            tipo_documento='LIQUIDACION', liquidacion=instance, estado='FIRMADO'
        ).exists()
        if firmada:
            return Response(
                {'error': 'Esta liquidación ya fue firmada por el trabajador y no puede modificarse.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        pendiente = SolicitudFirma.objects.filter(
            tipo_documento='LIQUIDACION', liquidacion=instance, estado='PENDIENTE'
        ).exists()
        if pendiente:
            return Response(
                {'error': 'Hay una solicitud de firma pendiente para esta liquidación. Cancélala antes de modificarla.'},
                status=status.HTTP_403_FORBIDDEN,
            )

        empleado = instance.empleado
        contrato = Contrato.objects.filter(empleado=empleado).first()

        if not contrato:
            return Response({'error': 'El trabajador no tiene un contrato activo.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            _validar_conceptos(request.data, request.user)
        except ValidationError as e:
            return Response(e.detail, status=status.HTTP_400_BAD_REQUEST)

        data = request.data
        campos_editables = [
            'mes', 'anio', 'dias_trabajados', 'dias_ausencia', 'dias_licencia',
            'dias_no_contratados', 'detalle_items',
        ]
        # Combina lo que venga en el request con lo que ya estaba guardado,
        # así un PATCH parcial recalcula usando el resto de los valores tal
        # como estaban, en vez de perderlos.
        datos_para_calculo = {
            campo: data.get(campo, getattr(instance, campo)) for campo in campos_editables
        }

        try:
            calculado = _calcular_liquidacion(
                contrato, empleado, datos_para_calculo,
                terminos=_terminos_congelados(instance, contrato),
            )

            for campo, valor in calculado.items():
                setattr(instance, campo, valor)

            instance.mes = data.get('mes', instance.mes)
            instance.anio = data.get('anio', instance.anio)

            # El PDF generado antes de este cambio ya no refleja los montos
            # recalculados — se limpia para forzar que se regenere.
            if instance.archivo_pdf:
                instance.archivo_pdf.delete(save=False)

            instance.save()
        except IntegrityError:
            return Response(
                {'error': 'Ya existe una liquidación para este trabajador en el período indicado.'},
                status=status.HTTP_409_CONFLICT
            )
        except Exception as e:
            return Response({'error': str(e)}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        serializer = self.get_serializer(instance)
        return Response(serializer.data)
        
    @action(detail=True, methods=['get'])
    def generar_pdf(self, request, pk=None):
        try:
            liquidacion = self.get_object()
            nombre_archivo = f'Liquidacion_{liquidacion.mes}_{liquidacion.anio}_{liquidacion.empleado.rut}.pdf'
            firmado = pdf_firmado('LIQUIDACION', liquidacion=liquidacion)
            if firmado:
                return respuesta_pdf(firmado, nombre_archivo, firmado=True)
            pdf = _pdf_liquidacion(liquidacion, _es_plan_semilla(request.user))
            if pdf is None:
                return Response({'error': 'Error al generar PDF'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
            return respuesta_pdf(pdf, nombre_archivo)
        except Exception as e:
            logger.exception('Error al generar PDF de liquidación')
            return Response({'error': f'Error generando PDF: {str(e)}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

    @action(detail=False, methods=['get'], url_path='zip_periodo')
    def zip_periodo(self, request):
        """PDF de todas las liquidaciones de una empresa en un período, en un ZIP (plan Pyme+)."""
        if not _plan_permite(request.user, 3):
            return Response(
                {'error': 'La descarga de liquidaciones en ZIP está disponible desde el plan Pyme. Mejora tu suscripción para acceder.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            mes = int(request.query_params.get('mes'))
            anio = int(request.query_params.get('anio'))
            empresa = Empresa.objects.get(id=request.query_params.get('empresa'), owner=request.user)
        except (TypeError, ValueError, Empresa.DoesNotExist):
            return Response({'error': 'Indica una empresa y un período válidos.'}, status=status.HTTP_400_BAD_REQUEST)

        liquidaciones = (Liquidacion.objects
                         .filter(empleado__empresa=empresa, mes=mes, anio=anio)
                         .select_related('empleado', 'empleado__empresa')
                         .order_by('empleado__apellido_paterno'))
        if not liquidaciones.exists():
            return Response({'error': 'No hay liquidaciones emitidas en ese período.'}, status=status.HTTP_404_NOT_FOUND)

        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            for liq in liquidaciones:
                # La firmada, si el trabajador la firmó; si no, la emitida.
                firmado = pdf_firmado('LIQUIDACION', liquidacion=liq)
                pdf = firmado or _pdf_liquidacion(liq, False)
                if pdf is None:
                    continue
                emp = liq.empleado
                sufijo = '_firmada' if firmado else ''
                nombre = f"Liquidacion_{emp.rut.replace('.', '')}_{emp.apellido_paterno}_{anio}-{mes:02d}{sufijo}.pdf".replace(' ', '_')
                zf.writestr(nombre, pdf)
        response = HttpResponse(buffer.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="Liquidaciones_{empresa.rut}_{anio}-{mes:02d}.zip"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'
        return response

    @action(detail=False, methods=['get'], url_path='exportar_previred')
    def exportar_previred(self, request):
        """Nómina del mes para Previred: formato estándar de largo variable, 105 campos.

        Si a algún trabajador le falta un dato sin el cual Previred rechaza el
        archivo, no se genera y se indica qué completar.
        """
        if not _plan_permite(request.user, 3):
            return Response(
                {'error': 'La exportación Previred está disponible desde el plan Pyme. Mejora tu suscripción para acceder.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            mes = int(request.query_params.get('mes'))
            anio = int(request.query_params.get('anio'))
        except (TypeError, ValueError):
            return Response({'error': 'Se requieren los parámetros mes y anio, numéricos.'}, status=400)

        qs = Liquidacion.objects.filter(
            empleado__empresa__owner=request.user, mes=mes, anio=anio,
        ).select_related('empleado', 'empleado__empresa', 'empleado__contrato_activo').order_by(
            'empleado__apellido_paterno', 'empleado__nombres')
        empresa_id = request.query_params.get('empresa')
        if empresa_id:
            qs = qs.filter(empleado__empresa_id=empresa_id)
        if not qs.exists():
            return Response({'error': 'No hay liquidaciones para el período seleccionado.'}, status=404)

        lineas, problemas = [], []
        for liq in qs:
            campos, faltan = _linea_previred(liq)
            if faltan:
                nombre = f'{liq.empleado.nombres} {liq.empleado.apellido_paterno}'.strip()
                problemas.append(f'{nombre}: {"; ".join(faltan)}')
            else:
                lineas.append(';'.join(campos))
        if problemas:
            return Response({'error': 'Completa estos datos para generar el archivo Previred (sin ellos Previred lo '
                                      'rechaza): ' + ' | '.join(problemas)}, status=400)

        meses = ['Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio', 'Julio', 'Agosto',
                 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre']
        nombre_archivo = f'Previred_{meses[mes - 1]}_{anio}.txt'
        # Previred lee el archivo en Latin-1 (ISO-8859-1), con fin de línea Windows.
        contenido = ('\r\n'.join(lineas) + '\r\n').encode('latin-1', errors='replace')
        response = HttpResponse(contenido, content_type='text/plain; charset=iso-8859-1')
        response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
        response['Access-Control-Expose-Headers'] = 'Content-Disposition'
        return response

    @action(detail=False, methods=['get'], url_path='libro_remuneraciones')
    def libro_remuneraciones(self, request):
        if not _plan_permite(request.user, 3):
            return Response(
                {'error': 'El Libro de Remuneraciones está disponible desde el plan Pyme. Mejora tu suscripción para acceder.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        mes_param  = request.query_params.get('mes')
        anio_param = request.query_params.get('anio')
        empresa_id = request.query_params.get('empresa')
        formato    = request.query_params.get('formato', 'excel')

        if not mes_param or not anio_param:
            return Response({'error': 'Se requieren los parámetros mes y anio.'}, status=400)
        try:
            mes  = int(mes_param)
            anio = int(anio_param)
        except ValueError:
            return Response({'error': 'Parámetros mes y anio deben ser numéricos.'}, status=400)

        qs = Liquidacion.objects.filter(
            empleado__empresa__owner=request.user,
            mes=mes, anio=anio,
        ).select_related('empleado', 'empleado__empresa').order_by(
            'empleado__ficha_numero', 'empleado__apellido_paterno'
        )
        if empresa_id:
            qs = qs.filter(empleado__empresa_id=empresa_id)
        if not qs.exists():
            return Response({'error': 'No hay liquidaciones para el período seleccionado.'}, status=404)

        empresa = qs.first().empleado.empresa
        meses_nombres = [
            'Enero', 'Febrero', 'Marzo', 'Abril', 'Mayo', 'Junio',
            'Julio', 'Agosto', 'Septiembre', 'Octubre', 'Noviembre', 'Diciembre',
        ]
        mes_nombre = meses_nombres[mes - 1]

        # ── Helper formato CLP para el PDF ────────────────────────────────────
        def clp(n):
            if not n:
                return '$0'
            return f'${int(n):,}'.replace(',', '.')

        # ── Preparación de filas (compartido por Excel y PDF) ─────────────────
        filas = []
        totales = {k: 0 for k in [
            'sueldo_base', 'gratificacion', 'otros_imp', 'total_imponible',
            'no_imponibles', 'total_haberes', 'cotiz_afp', 'cotiz_salud',
            'cesantia', 'imp_unico', 'anticipo', 'otros_desc',
            'total_descuentos', 'sueldo_liquido',
        ]}

        for liq in qs:
            emp = liq.empleado
            agrupados = liq.items_agrupados
            det_imp = agrupados['imponibles']
            det_hex = agrupados['horas_extras']
            det_com = agrupados['comisiones']
            det_noi = agrupados['no_imponibles']
            det_odc = agrupados['descuentos']

            otros_imp  = sum(int(d.get('valor', 0)) for d in det_imp if isinstance(d, dict))
            otros_imp += sum(int(d.get('valor', 0)) for d in det_hex if isinstance(d, dict))
            otros_imp += sum(int(d.get('valor', 0)) for d in det_com if isinstance(d, dict))
            otros_imp += int(liq.semana_corrida or 0)
            no_impon   = sum(int(d.get('valor', 0)) for d in det_noi if isinstance(d, dict))
            otros_desc = sum(int(d.get('valor', 0)) for d in det_odc if isinstance(d, dict))

            filas.append({
                'ficha':        emp.ficha_numero or '',
                'rut':          emp.rut,
                'nombre':       f'{emp.apellido_paterno} {emp.apellido_materno or ""} {emp.nombres}'.strip(),
                'cargo':        emp.cargo or '',
                'dias':         liq.dias_trabajados,
                'sueldo_base':  int(liq.sueldo_base),
                'gratificacion':int(liq.gratificacion),
                'otros_imp':    otros_imp,
                'total_imp':    int(liq.total_imponible),
                'no_imp':       no_impon,
                'total_hab':    int(liq.total_haberes),
                'afp_nombre':   liq.afp_nombre or '',
                'cotiz_afp':    int(liq.afp_monto),
                'salud_nombre': liq.salud_nombre or '',
                'cotiz_salud':  int(liq.salud_monto),
                'cesantia':     int(liq.seguro_cesantia),
                'imp_unico':    int(liq.impuesto_unico or 0),
                'anticipo':     int(liq.anticipo_quincena or 0),
                'otros_desc':   otros_desc,
                'total_desc':   int(liq.total_descuentos),
                'sueldo_liq':   int(liq.sueldo_liquido),
            })

            totales['sueldo_base']      += int(liq.sueldo_base)
            totales['gratificacion']    += int(liq.gratificacion)
            totales['otros_imp']        += otros_imp
            totales['total_imponible']  += int(liq.total_imponible)
            totales['no_imponibles']    += no_impon
            totales['total_haberes']    += int(liq.total_haberes)
            totales['cotiz_afp']        += int(liq.afp_monto)
            totales['cotiz_salud']      += int(liq.salud_monto)
            totales['cesantia']         += int(liq.seguro_cesantia)
            totales['imp_unico']        += int(liq.impuesto_unico or 0)
            totales['anticipo']         += int(liq.anticipo_quincena or 0)
            totales['otros_desc']       += otros_desc
            totales['total_descuentos'] += int(liq.total_descuentos)
            totales['sueldo_liquido']   += int(liq.sueldo_liquido)

        nombre_base = f'LibroRemuneraciones_{mes_nombre}_{anio}_{empresa.rut}'

        # ══════════════════════════════════════════════════════════════════════
        # RAMA PDF
        # ══════════════════════════════════════════════════════════════════════
        if formato == 'pdf':
            # Pre-formatear montos para el template
            for f in filas:
                for k in ['sueldo_base','gratificacion','otros_imp','total_imp','no_imp',
                          'total_hab','cotiz_afp','cotiz_salud','cesantia','imp_unico',
                          'anticipo','otros_desc','total_desc','sueldo_liq']:
                    f[f'{k}_fmt'] = clp(f[k])

            totales_fmt = {k: clp(v) for k, v in totales.items()}

            context = {
                'empresa':      empresa,
                'mes_nombre':   mes_nombre,
                'anio':         anio,
                'filas':        filas,
                'totales':      totales,
                'totales_fmt':  totales_fmt,
                'fecha_emision': timezone.localdate().strftime('%d/%m/%Y'),
                'n_trabajadores': len(filas),
            }
            template = get_template('libro_remuneraciones.html')
            html = template.render(context)

            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{nombre_base}.pdf"'
            pisa_status = pisa.CreatePDF(html, dest=response)
            if pisa_status.err:
                return Response({'error': 'Error al generar PDF'}, status=500)
            return response

        # ══════════════════════════════════════════════════════════════════════
        # RAMA EXCEL (sin cambios respecto al paso 1)
        # ══════════════════════════════════════════════════════════════════════

        # ── Estilos ───────────────────────────────────────────────────────────
        COLOR_HEADER   = '1E3A5F'
        COLOR_TOTALES  = 'F59E0B'
        COLOR_FILA_PAR = 'F1F5F9'

        ft_titulo    = Font(name='Calibri', bold=True, size=14, color='1E3A5F')
        ft_subtit    = Font(name='Calibri', bold=True, size=11, color='334155')
        ft_header    = Font(name='Calibri', bold=True, size=9,  color='FFFFFF')
        ft_dato      = Font(name='Calibri', size=9)
        ft_total     = Font(name='Calibri', bold=True, size=9)
        ft_total_liq = Font(name='Calibri', bold=True, size=9, color='7C3AED')

        al_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        al_left   = Alignment(horizontal='left',   vertical='center')
        al_right  = Alignment(horizontal='right',  vertical='center')

        fill_header = PatternFill('solid', fgColor=COLOR_HEADER)
        fill_total  = PatternFill('solid', fgColor=COLOR_TOTALES)
        fill_par    = PatternFill('solid', fgColor=COLOR_FILA_PAR)

        thin  = Side(style='thin', color='CBD5E1')
        borde = Border(left=thin, right=thin, top=thin, bottom=thin)

        FMT_CLP  = '#,##0'
        FMT_TEXT = '@'

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f'Libro {mes_nombre} {anio}'
        ws.sheet_view.showGridLines = False

        COLS = [
            ('N°',                    5),
            ('RUT',                  12),
            ('Apellidos y Nombres',  28),
            ('Cargo',                16),
            ('Días\nTrab.',           7),
            ('Sueldo\nBase',         12),
            ('Gratif.',              11),
            ('Otros Hab.\nImpon.',   12),
            ('Total\nImponible',     13),
            ('Hab. No\nImponibles',  13),
            ('Total\nHaberes',       12),
            ('AFP',                   9),
            ('Cotiz.\nAFP',          11),
            ('Salud',                 9),
            ('Cotiz.\nSalud',        11),
            ('Cesantía',             10),
            ('Imp.\nÚnico',          10),
            ('Anticipo',             10),
            ('Otros\nDesc.',         10),
            ('Total\nDescuentos',    13),
            ('Alcance\nLíquido',     13),
        ]
        for i, (_, ancho) in enumerate(COLS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = ancho

        N_COLS     = len(COLS)
        ultima_col = get_column_letter(N_COLS)

        ws.row_dimensions[1].height = 22
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 16

        ws.merge_cells(f'A1:{ultima_col}1')
        c = ws['A1']
        c.value     = empresa.nombre_legal.upper()
        c.font      = ft_titulo
        c.alignment = al_center

        ws.merge_cells(f'A2:{ultima_col}2')
        c = ws['A2']
        c.value     = 'LIBRO DE REMUNERACIONES'
        c.font      = ft_subtit
        c.alignment = al_center

        ws.merge_cells(f'A3:{ultima_col}3')
        c = ws['A3']
        c.value     = f'RUT: {empresa.rut}     Período: {mes_nombre} {anio}'
        c.font      = Font(name='Calibri', size=10, color='475569')
        c.alignment = al_center

        ws.row_dimensions[4].height = 6

        ws.row_dimensions[5].height = 36
        for col_idx, (label, _) in enumerate(COLS, start=1):
            c = ws.cell(row=5, column=col_idx, value=label)
            c.font      = ft_header
            c.fill      = fill_header
            c.alignment = al_center
            c.border    = borde

        for fila_idx, f in enumerate(filas, start=6):
            fill_fila = fill_par if fila_idx % 2 == 0 else None
            valores = [
                (f['ficha'],         FMT_TEXT, al_center),
                (f['rut'],           FMT_TEXT, al_left),
                (f['nombre'],        FMT_TEXT, al_left),
                (f['cargo'],         FMT_TEXT, al_left),
                (f['dias'],          FMT_TEXT, al_center),
                (f['sueldo_base'],   FMT_CLP,  al_right),
                (f['gratificacion'], FMT_CLP,  al_right),
                (f['otros_imp'],     FMT_CLP,  al_right),
                (f['total_imp'],     FMT_CLP,  al_right),
                (f['no_imp'],        FMT_CLP,  al_right),
                (f['total_hab'],     FMT_CLP,  al_right),
                (f['afp_nombre'],    FMT_TEXT, al_center),
                (f['cotiz_afp'],     FMT_CLP,  al_right),
                (f['salud_nombre'],  FMT_TEXT, al_center),
                (f['cotiz_salud'],   FMT_CLP,  al_right),
                (f['cesantia'],      FMT_CLP,  al_right),
                (f['imp_unico'],     FMT_CLP,  al_right),
                (f['anticipo'],      FMT_CLP,  al_right),
                (f['otros_desc'],    FMT_CLP,  al_right),
                (f['total_desc'],    FMT_CLP,  al_right),
                (f['sueldo_liq'],    FMT_CLP,  al_right),
            ]
            ws.row_dimensions[fila_idx].height = 15
            for col_idx, (valor, fmt, alin) in enumerate(valores, start=1):
                c = ws.cell(row=fila_idx, column=col_idx, value=valor)
                c.number_format = fmt
                c.font          = ft_dato
                c.alignment     = alin
                c.border        = borde
                if fill_fila:
                    c.fill = fill_fila

        fila_total = len(filas) + 6
        ws.row_dimensions[fila_total].height = 18
        vals_total = [
            ('', FMT_TEXT), ('', FMT_TEXT),
            ('TOTALES', FMT_TEXT), ('', FMT_TEXT), ('', FMT_TEXT),
            (totales['sueldo_base'],      FMT_CLP),
            (totales['gratificacion'],    FMT_CLP),
            (totales['otros_imp'],        FMT_CLP),
            (totales['total_imponible'],  FMT_CLP),
            (totales['no_imponibles'],    FMT_CLP),
            (totales['total_haberes'],    FMT_CLP),
            ('', FMT_TEXT),
            (totales['cotiz_afp'],        FMT_CLP),
            ('', FMT_TEXT),
            (totales['cotiz_salud'],      FMT_CLP),
            (totales['cesantia'],         FMT_CLP),
            (totales['imp_unico'],        FMT_CLP),
            (totales['anticipo'],         FMT_CLP),
            (totales['otros_desc'],       FMT_CLP),
            (totales['total_descuentos'], FMT_CLP),
            (totales['sueldo_liquido'],   FMT_CLP),
        ]
        for col_idx, (valor, fmt) in enumerate(vals_total, start=1):
            c = ws.cell(row=fila_total, column=col_idx, value=valor)
            c.number_format = fmt
            c.fill      = fill_total
            c.border    = borde
            c.alignment = al_right if fmt == FMT_CLP else al_center
            if col_idx == 3:
                c.alignment = al_left
            c.font = ft_total_liq if col_idx == N_COLS else ft_total

        ws.page_setup.orientation   = 'landscape'
        ws.page_setup.paperSize     = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToPage     = True
        ws.page_setup.fitToWidth    = 1
        ws.page_setup.fitToHeight   = 0
        ws.print_title_rows         = '1:5'
        ws.freeze_panes             = 'A6'

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{nombre_base}.xlsx"'
        return response

    # ──────────────────────────────────────────────────────────────────────────
    @action(detail=False, methods=['get'], url_path='consolidado')
    def consolidado(self, request):
        # Consolidado multiempresa: plan Pyme en adelante.
        if not _plan_permite(request.user, 3):
            return Response(
                {'error': 'El consolidado multiempresa está disponible desde el plan Pyme. Mejora tu suscripción para acceder.'},
                status=status.HTTP_403_FORBIDDEN,
            )
        anio_param = request.query_params.get('anio')
        mes_param  = request.query_params.get('mes')
        formato    = request.query_params.get('formato', 'json')

        if not anio_param:
            return Response({'error': 'Se requiere el parámetro anio.'}, status=400)
        try:
            anio = int(anio_param)
            mes  = int(mes_param) if mes_param else None
        except ValueError:
            return Response({'error': 'Los parámetros anio y mes deben ser numéricos.'}, status=400)

        base_qs = (
            Liquidacion.objects
            .filter(empleado__empresa__owner=request.user, anio=anio)
            .select_related('empleado', 'empleado__empresa', 'empleado__contrato_activo')
        )
        qs_periodo = base_qs.filter(mes=mes) if mes else base_qs

        if not qs_periodo.exists():
            if formato == 'json':
                return Response({'error': 'No hay liquidaciones para el período seleccionado.'}, status=404)
            return Response({'error': 'No hay liquidaciones para el período seleccionado.'}, status=404)

        def _costo_emp(liq):
            """Aportes de cargo del empleador: SIS, mutual, cesantía y reforma."""
            par = _parametros_previsionales(liq.mes, liq.anio)
            contrato_liq = getattr(liq.empleado, 'contrato_activo', None)
            tipo = liq.tipo_contrato or getattr(contrato_liq, 'tipo_contrato', 'INDEFINIDO')
            _, tasa_afc = _tasas_afc(par, tipo, _anios_de_servicio(liq.empleado, contrato_liq, liq.mes, liq.anio))
            uf_periodo = float(liq.valor_uf) or obtener_uf()
            renta = min(
                int(liq.total_imponible or 0),
                _tope_en_pesos(par['tope_imponible_afp_uf'], uf_periodo),
            )
            renta_afc = min(
                int(liq.total_imponible or 0),
                _tope_en_pesos(par['tope_imponible_afc_uf'], uf_periodo),
            )
            # Reforma de pensiones (Ley 21.735): aporte a la cuenta individual,
            # rentabilidad protegida y expectativa de vida.
            reforma = par['tasa_afp_empleador'] + par['tasa_rentabilidad_protegida'] + par['tasa_expectativa_vida']
            return int(renta * (par['tasa_sis'] + _tasa_accidentes(liq.empleado.empresa, par) + reforma) + renta_afc * tasa_afc)

        def clp(n):
            if not n: return '$0'
            return f'${int(n):,}'.replace(',', '.')

        # ── Datos compartidos ─────────────────────────────────────────────────
        from collections import defaultdict
        masa_salarial = liquido_total = costo_empleador = 0
        empleados_ids = set()
        empresas_dict = defaultdict(lambda: {
            'id': None, 'nombre': '', 'rut': '',
            'trabajadores': 0, 'masa_salarial': 0,
            'liquido_total': 0, 'costo_empleador': 0,
        })
        for liq in qs_periodo:
            ce = _costo_emp(liq)
            masa_salarial   += liq.total_haberes
            liquido_total   += liq.sueldo_liquido
            costo_empleador += ce
            empleados_ids.add(liq.empleado_id)
            eid = liq.empleado.empresa_id
            emp = liq.empleado.empresa
            d = empresas_dict[eid]
            d['id']             = eid
            d['nombre']         = emp.nombre_legal
            d['rut']            = emp.rut
            d['trabajadores']  += 1
            d['masa_salarial'] += liq.total_haberes
            d['liquido_total'] += liq.sueldo_liquido
            d['costo_empleador'] += ce

        empresas_list = sorted(empresas_dict.values(), key=lambda x: x['masa_salarial'], reverse=True)

        MESES_CORTOS = ['Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Sep','Oct','Nov','Dic']
        MESES_LARGOS = ['Enero','Febrero','Marzo','Abril','Mayo','Junio',
                        'Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']
        evolucion = []
        for m in range(1, 13):
            qs_m = base_qs.filter(mes=m)
            if qs_m.exists():
                ms = sum(l.total_haberes   for l in qs_m)
                lq = sum(l.sueldo_liquido  for l in qs_m)
                ce = sum(_costo_emp(l)     for l in qs_m)
                tw = qs_m.values('empleado_id').distinct().count()
            else:
                ms = lq = ce = tw = 0
            evolucion.append({'mes': m, 'mes_nombre': MESES_CORTOS[m-1],
                              'masa_salarial': ms, 'liquido_total': lq,
                              'costo_empleador': ce, 'trabajadores': tw})

        kpis = {
            'masa_salarial':   masa_salarial,
            'trabajadores':    len(empleados_ids),
            'costo_empleador': costo_empleador,
            'liquido_total':   liquido_total,
        }
        periodo_nombre = MESES_LARGOS[mes - 1] if mes else f'Año {anio}'
        nombre_base = f'Consolidado_{periodo_nombre.replace(" ","_")}_{anio}'

        # ══════════════════════════════════════════════════════════════════════
        # RAMA JSON
        # ══════════════════════════════════════════════════════════════════════
        if formato == 'json':
            return Response({
                'periodo': {'anio': anio, 'mes': mes, 'mes_nombre': MESES_LARGOS[mes-1] if mes else None},
                'kpis': kpis,
                'empresas': empresas_list,
                'evolucion': evolucion,
            })

        # ══════════════════════════════════════════════════════════════════════
        # RAMA EXCEL
        # ══════════════════════════════════════════════════════════════════════
        if formato == 'excel':
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter

            wb = Workbook()

            # ── Hoja 1: Resumen ───────────────────────────────────────────────
            ws = wb.active
            ws.title = 'Resumen'

            hdr_fill  = PatternFill('solid', fgColor='1E3A5F')
            tot_fill  = PatternFill('solid', fgColor='FEF3C7')
            kpi_fill  = PatternFill('solid', fgColor='0F2540')
            thin = Side(style='thin', color='CCCCCC')
            bord = Border(left=thin, right=thin, top=thin, bottom=thin)

            def _set(cell, val, bold=False, fill=None, align='left', color='000000'):
                cell.value = val
                cell.font  = Font(bold=bold, color=color, size=10)
                cell.alignment = Alignment(horizontal=align, vertical='center', wrap_text=True)
                if fill: cell.fill = fill
                cell.border = bord

            # Title
            ws.merge_cells('A1:G1')
            t = ws['A1']
            t.value = f'Reporte Consolidado de Remuneraciones · {periodo_nombre} {anio}'
            t.font  = Font(bold=True, size=13, color='FFFFFF')
            t.fill  = PatternFill('solid', fgColor='0C1A35')
            t.alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[1].height = 28

            ws.merge_cells('A2:G2')
            ws['A2'].value = f'Generado el {timezone.localdate().strftime("%d/%m/%Y")} · {len(empresas_list)} empresa(s) · {kpis["trabajadores"]} trabajadore(s)'
            ws['A2'].font  = Font(size=9, color='888888')
            ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
            ws.row_dimensions[2].height = 16

            # KPI row
            kpi_labels = ['Masa salarial', 'Trabajadores', 'Costo empleador', 'Líquido a pagar']
            kpi_vals   = [clp(kpis['masa_salarial']), str(kpis['trabajadores']),
                          clp(kpis['costo_empleador']), clp(kpis['liquido_total'])]
            kpi_cols   = [('A','B'), ('C','C'), ('D','E'), ('F','G')]
            ws.row_dimensions[3].height = 14
            ws.row_dimensions[4].height = 22
            ws.row_dimensions[5].height = 22
            ws.row_dimensions[6].height = 8
            for (c1, c2), lbl, val in zip(kpi_cols, kpi_labels, kpi_vals):
                ws.merge_cells(f'{c1}4:{c2}4')
                ws.merge_cells(f'{c1}5:{c2}5')
                lc = ws[f'{c1}4']
                lc.value = lbl; lc.font = Font(bold=True, size=8, color='AAAAAA')
                lc.fill = kpi_fill; lc.alignment = Alignment(horizontal='center', vertical='center')
                vc = ws[f'{c1}5']
                vc.value = val; vc.font = Font(bold=True, size=11, color='FFFFFF')
                vc.fill = kpi_fill; vc.alignment = Alignment(horizontal='center', vertical='center')

            # Table header
            cols_h = ['Empresa', 'RUT', 'Trabajadores', 'Masa Salarial', 'Costo Empleador', 'Líquido', '% del Total']
            for ci, h in enumerate(cols_h, 1):
                c = ws.cell(row=7, column=ci)
                _set(c, h, bold=True, fill=hdr_fill, align='center', color='FFFFFF')
            ws.row_dimensions[7].height = 20

            for ri, emp in enumerate(empresas_list, 8):
                pct = round(emp['masa_salarial'] / masa_salarial * 100, 1) if masa_salarial else 0
                row_fill = PatternFill('solid', fgColor='F0F4F8') if ri % 2 == 0 else None
                vals = [emp['nombre'], emp['rut'], emp['trabajadores'],
                        clp(emp['masa_salarial']), clp(emp['costo_empleador']),
                        clp(emp['liquido_total']), f'{pct}%']
                aligns = ['left','center','center','right','right','right','center']
                for ci, (v, a) in enumerate(zip(vals, aligns), 1):
                    _set(ws.cell(row=ri, column=ci), v, fill=row_fill, align=a)
                ws.row_dimensions[ri].height = 18

            # Totals
            tr = len(empresas_list) + 8
            tot_vals = ['TOTAL CONSOLIDADO', '', kpis['trabajadores'],
                        clp(kpis['masa_salarial']), clp(kpis['costo_empleador']),
                        clp(kpis['liquido_total']), '100%']
            for ci, v in enumerate(tot_vals, 1):
                a = 'right' if ci >= 4 else ('center' if ci == 3 else 'left')
                _set(ws.cell(row=tr, column=ci), v, bold=True, fill=tot_fill, align=a)
            ws.row_dimensions[tr].height = 20

            # Column widths
            for ci, w in enumerate([38, 14, 13, 18, 18, 18, 12], 1):
                ws.column_dimensions[get_column_letter(ci)].width = w
            ws.freeze_panes = 'A8'

            # ── Hoja 2: Evolución mensual ─────────────────────────────────────
            ws2 = wb.create_sheet('Evolución mensual')
            evo_headers = ['Mes', 'Masa Salarial', 'Líquido', 'Costo Empleador', 'Trabajadores']
            for ci, h in enumerate(evo_headers, 1):
                c = ws2.cell(row=1, column=ci)
                _set(c, h, bold=True, fill=hdr_fill, align='center', color='FFFFFF')
                ws2.row_dimensions[1].height = 20

            for ri, ev in enumerate(evolucion, 2):
                row_fill = PatternFill('solid', fgColor='F0F4F8') if ri % 2 == 0 else None
                vals = [ev['mes_nombre'], clp(ev['masa_salarial']), clp(ev['liquido_total']),
                        clp(ev['costo_empleador']), ev['trabajadores']]
                aligns = ['center','right','right','right','center']
                for ci, (v, a) in enumerate(zip(vals, aligns), 1):
                    _set(ws2.cell(row=ri, column=ci), v, fill=row_fill, align=a)
                ws2.row_dimensions[ri].height = 17

            for ci, w in enumerate([14, 18, 18, 18, 14], 1):
                ws2.column_dimensions[get_column_letter(ci)].width = w

            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize   = ws.PAPERSIZE_LETTER
            ws2.page_setup.orientation = 'landscape'
            ws2.page_setup.paperSize   = ws2.PAPERSIZE_LETTER

            buf = io.BytesIO()
            wb.save(buf); buf.seek(0)
            response = HttpResponse(buf.read(),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{nombre_base}.xlsx"'
            return response

        # ══════════════════════════════════════════════════════════════════════
        # RAMA PDF
        # ══════════════════════════════════════════════════════════════════════
        if formato == 'pdf':
            for emp in empresas_list:
                emp['pct'] = round(emp['masa_salarial'] / masa_salarial * 100, 1) if masa_salarial else 0
                emp['masa_salarial_fmt']   = clp(emp['masa_salarial'])
                emp['costo_empleador_fmt'] = clp(emp['costo_empleador'])
                emp['liquido_total_fmt']   = clp(emp['liquido_total'])
            for ev in evolucion:
                ev['masa_salarial_fmt']   = clp(ev['masa_salarial'])
                ev['liquido_total_fmt']   = clp(ev['liquido_total'])
                ev['costo_empleador_fmt'] = clp(ev['costo_empleador'])

            context = {
                'periodo_nombre':  periodo_nombre,
                'anio':            anio,
                'mes':             mes,
                'kpis':            kpis,
                'kpis_fmt': {
                    'masa_salarial':   clp(kpis['masa_salarial']),
                    'costo_empleador': clp(kpis['costo_empleador']),
                    'liquido_total':   clp(kpis['liquido_total']),
                    'trabajadores':    kpis['trabajadores'],
                },
                'empresas':        empresas_list,
                'evolucion':       evolucion,
                'n_empresas':      len(empresas_list),
                'fecha_emision':   timezone.localdate().strftime('%d/%m/%Y'),
            }
            template = get_template('consolidado_remuneraciones.html')
            html = template.render(context)
            response = HttpResponse(content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="{nombre_base}.pdf"'
            pisa_status = pisa.CreatePDF(html, dest=response)
            if pisa_status.err:
                return Response({'error': 'Error al generar PDF'}, status=500)
            return response

        return Response({'error': 'Formato no válido. Use json, excel o pdf.'}, status=400)
